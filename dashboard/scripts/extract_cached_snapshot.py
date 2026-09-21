"""Extract deployment-friendly CSV snapshots from cached workbook pivots.

The workbook is opened read-only and is never saved or modified. The source
ranges below are the cached Pivot Support blocks documented in this project.
"""

from __future__ import annotations

import csv
import json
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
WORKBOOK = ROOT / "workbook" / "London_Transport_Demand_Analytics_Portfolio.xlsx"
OUTPUT = ROOT / "dashboard" / "data" / "cached"


def values(ws, cell_range: str) -> list[list[object]]:
    return [[cell.value for cell in row] for row in ws[cell_range]]


def write_csv(name: str, headers: list[str], rows: list[list[object]]) -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    with (OUTPUT / name).open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(headers)
        writer.writerows(rows)


def non_total(rows: list[list[object]]) -> list[list[object]]:
    return [row for row in rows if row and row[0] not in (None, "", "Grand Total", "Row Labels")]


def main() -> None:
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    pivot = workbook["Pivot Support"]

    write_csv(
        "network_monthly.csv",
        ["month", "tube_avg_daily", "bus_avg_daily"],
        non_total(values(pivot, "H57:J149")),
    )
    write_csv(
        "network_annual.csv",
        ["year", "tube_avg_daily", "bus_avg_daily"],
        non_total(values(pivot, "H43:J51")),
    )
    write_csv(
        "network_weekday.csv",
        ["day_of_week", "bus_avg_daily", "tube_avg_daily"],
        non_total(values(pivot, "H16:J23")),
    )
    write_csv(
        "top_stations.csv",
        ["station", "average_daily_footfall"],
        non_total(values(pivot, "L57:M67")),
    )
    write_csv(
        "station_monthly_detail.csv",
        ["month", "average_daily_footfall"],
        non_total(values(pivot, "A4:B11")),
    )
    write_csv(
        "station_weekday_detail.csv",
        ["day_of_week", "average_daily_footfall"],
        non_total(values(pivot, "D6:E11")),
    )
    write_csv(
        "station_entries_exits_detail.csv",
        ["month", "average_daily_entries", "average_daily_exits"],
        non_total(values(pivot, "A99:C106")),
    )
    write_csv(
        "service_reporting_period.csv",
        ["reporting_period", "service_delivery", "peak_delivery", "off_peak_delivery"],
        non_total(values(pivot, "O5:R18")),
    )
    write_csv(
        "service_by_line.csv",
        ["line", "service_delivery"],
        non_total(values(pivot, "O21:P31")),
    )
    write_csv(
        "service_peak_offpeak_by_line.csv",
        ["line", "peak_delivery", "off_peak_delivery"],
        non_total(values(pivot, "R21:T31")),
    )
    write_csv(
        "executive_service_by_line.csv",
        ["line", "service_delivery"],
        non_total(values(pivot, "L71:M81")),
    )
    write_csv(
        "numbat_gate_evening.csv",
        ["quarter_hour", "entries", "exits"],
        non_total(values(pivot, "O38:Q50")),
    )
    write_csv(
        "numbat_platform_evening.csv",
        ["quarter_hour", "boarders", "alighters"],
        non_total(values(pivot, "S38:U50")),
    )
    write_csv(
        "numbat_intensity_evening.csv",
        ["link", "passengers_per_scheduled_train"],
        non_total(values(pivot, "W38:X48")),
    )
    write_csv(
        "numbat_link_load_evening.csv",
        ["link", "passenger_load"],
        non_total(values(pivot, "Z38:AA48")),
    )
    write_csv(
        "executive_numbat_intensity.csv",
        ["link", "passengers_per_scheduled_train"],
        non_total(values(pivot, "L84:M94")),
    )

    source_rows = values(workbook["Source Register"], "A5:I8")
    write_csv(
        "source_register.csv",
        [
            "source_id",
            "dataset",
            "provider",
            "file",
            "coverage",
            "refresh_frequency",
            "purpose",
            "status",
            "notes",
        ],
        source_rows,
    )
    quality_rows = values(workbook["Data Quality"], "A5:C14")
    write_csv("data_quality.csv", ["check", "result", "status"], quality_rows)

    station_kpi = values(pivot, "A1:G2")
    network_kpi = values(pivot, "I1:M2")
    service_kpi = values(pivot, "O1:S2")
    executive_kpi = values(pivot, "H53:L54")
    numbat_kpi = values(pivot, "O34:T35")

    def serialisable(value: object) -> object:
        return value.isoformat() if hasattr(value, "isoformat") else value

    def pair(block: list[list[object]]) -> dict[str, object]:
        return {
            str(key): serialisable(value)
            for key, value in zip(block[0], block[1], strict=False)
        }

    metadata = {
        "snapshot_version": 1,
        "source_workbook": WORKBOOK.name,
        "portfolio_build_refresh": "2026-08-17T15:38:00",
        "snapshot_extracted_at": datetime.now(timezone.utc).isoformat(),
        "journey_data_through": "2026-08-08",
        "station_data_through": "2026-07-04",
        "station_detail_context": {
            "station": "Langdon Park",
            "calendar_year": 2026,
            "day_type": "Weekday",
        },
        "service_detail_context": "Cached workbook selection across 2016/17–2023/24",
        "executive_service_context": "2025/26 cached executive selection",
        "numbat_context": "2024 TWT typical day; cached Evening band (19:00–22:00)",
        "kpis": {
            "station_detail": pair(station_kpi),
            "network_all_periods": pair(network_kpi),
            "service_detail": pair(service_kpi),
            "executive": pair(executive_kpi),
            "numbat_cached": pair(numbat_kpi),
        },
        "limitations": [
            "Missing station-day rows are not converted to zero.",
            "Station Footfall and Network Journey source dates differ.",
            "The cached workbook exposes pivot outputs, not the underlying million-row model tables.",
            "Station detail is available for the cached Langdon Park / 2026 / Weekday context only.",
            "Service performance uses the workbook's cached financial-year selection.",
            "NUMBAT is a typical 2024 Tue/Wed/Thu profile, not annual daily demand.",
            "Passengers per Scheduled Train is demand intensity, not occupancy or capacity utilisation.",
        ],
    }
    with (OUTPUT / "snapshot_meta.json").open("w", encoding="utf-8") as handle:
        json.dump(metadata, handle, indent=2, ensure_ascii=False)

    workbook.close()
    print(f"Wrote cached dashboard snapshots to {OUTPUT}")


if __name__ == "__main__":
    main()
